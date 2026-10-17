"""Write anything: Word, Excel (formulas + charts), PowerPoint, PDF (merge, split, fill forms, sign-ready),
Markdown -> a good-looking PDF (session 29: documents & data).

    write_docx(path, markdown, title)                 python-docx
    edit_docx(path, replace={old: new}, append=md)    python-docx (run formatting kept where the text sits in one run)
    write_xlsx(path, sheets=[...], charts=[...])      openpyxl - "=SUM(B2:B9)" cells are real formulas
    edit_xlsx(path, cells={"Sheet1!B2": 5})           openpyxl
    write_pptx(path, slides=[...], title)             python-pptx
    pdf_merge / pdf_split / pdf_fields / pdf_fill / pdf_sign_ready / pdf_info      pypdf
    markdown_to_pdf(markdown, path, title)            Chromium (the chromium add-on) > reportlab > LOKI's own
                                                      small PDF writer (always works, standard library)

python-docx, openpyxl, python-pptx and pypdf are in engine/requirements.txt (pure Python, small); every
function says plainly which package is missing when one is. Paths are checked by the caller (loki.py keeps
them inside the workspace). Standard library at import time.
"""
from __future__ import annotations

import datetime
import io
import os
import re
import zlib
from pathlib import Path

MISSING = "%s is not installed - run setup again or: pip install %s (it is in engine/requirements.txt)"


class WriteError(Exception):
    pass


def _need(module: str, pip_name: str):
    import importlib  # pylint: disable=import-outside-toplevel
    try:
        return importlib.import_module(module)
    except ImportError as exc:
        raise WriteError(MISSING % (pip_name, pip_name)) from exc


# --------------------------------------------------------------------------- Markdown -> blocks
def inline_runs(text: str) -> list:
    """[(text, bold, italic, code)] from **bold**, *italic*/_italic_, `code`; [label](url) keeps the label."""
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", str(text or ""))
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    out = []
    pattern = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False, False))
        tok = m.group(0)
        if tok.startswith(("**", "__")):
            out.append((tok[2:-2], True, False, False))
        elif tok.startswith("`"):
            out.append((tok[1:-1], False, False, True))
        else:
            out.append((tok[1:-1], False, True, False))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False, False))
    return [r for r in out if r[0]]


def plain(text: str) -> str:
    return "".join(r[0] for r in inline_runs(text))


def blocks(markdown: str) -> list:
    """A small, forgiving Markdown parser: [{"type": heading|para|bullet|number|table|code|quote|hr|image|pagebreak,
    ...}]. Headings carry level; lists carry level (indent // 2); tables carry rows; images carry path + alt."""
    lines = str(markdown or "").replace("\r\n", "\n").split("\n")
    out, i, para = [], 0, []

    def flush():
        if para:
            out.append({"type": "para", "text": " ".join(s.strip() for s in para)})
            para.clear()
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            lang, code = s[3:].strip(), []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            out.append({"type": "code", "text": "\n".join(code), "lang": lang})
            i += 1
            continue
        if not s:
            flush()
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*(\{#[^}]*\})?\s*#*$", s)
        if m:
            flush()
            out.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2)})
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush()
            out.append({"type": "hr"})
            i += 1
            continue
        if s in ("\\pagebreak", "<!-- pagebreak -->", "\\newpage"):
            flush()
            out.append({"type": "pagebreak"})
            i += 1
            continue
        img = re.match(r"^!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)$", s)
        if img:
            flush()
            out.append({"type": "image", "alt": img.group(1), "path": img.group(2)})
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{2,}", lines[i + 1].strip()):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                if not re.match(r"^\|?[\s:|-]+$", row):
                    cells = re.split(r"(?<!\\)\|", row.strip("|"))
                    rows.append([c.strip().replace("\\|", "|") for c in cells])
                i += 1
            out.append({"type": "table", "rows": rows})
            continue
        lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if lm:
            flush()
            kind = "number" if lm.group(2)[0].isdigit() else "bullet"
            text = lm.group(3)
            check = re.match(r"^\[( |x|X)\]\s+(.*)$", text)
            if check:
                text = ("☑ " if check.group(1).lower() == "x" else "☐ ") + check.group(2)
            out.append({"type": kind, "level": min(4, len(lm.group(1).replace("\t", "  ")) // 2), "text": text})
            i += 1
            continue
        if s.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            out.append({"type": "quote", "text": " ".join(quote)})
            continue
        para.append(line)
        i += 1
    flush()
    return out


# --------------------------------------------------------------------------- DOCX
def _docx_runs(paragraph, text: str, size=None) -> None:
    for t, bold, italic, code in inline_runs(text):
        run = paragraph.add_run(t)
        run.bold = bold or None
        run.italic = italic or None
        if code:
            run.font.name = "Consolas"
        if size:
            from docx.shared import Pt  # pylint: disable=import-outside-toplevel
            run.font.size = Pt(size)


def _docx_add_blocks(doc, items: list, base: Path | None) -> None:
    from docx.shared import Inches, Pt  # pylint: disable=import-outside-toplevel
    from docx.enum.text import WD_BREAK  # pylint: disable=import-outside-toplevel
    for b in items:
        kind = b["type"]
        if kind == "heading":
            doc.add_heading(plain(b["text"]), level=min(9, b["level"]))
        elif kind == "para":
            _docx_runs(doc.add_paragraph(), b["text"])
        elif kind in ("bullet", "number"):
            style = ("List Bullet" if kind == "bullet" else "List Number") + ("" if not b["level"] else " %d" % min(3, b["level"] + 1))
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet" if kind == "bullet" else "List Number")
            _docx_runs(p, b["text"])
        elif kind == "quote":
            try:
                p = doc.add_paragraph(style="Quote")
            except KeyError:
                p = doc.add_paragraph()
            _docx_runs(p, b["text"])
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b["text"])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif kind == "table" and b["rows"]:
            width = max(len(r) for r in b["rows"])
            table = doc.add_table(rows=0, cols=width)
            try:
                table.style = "Table Grid"
            except (KeyError, ValueError):
                pass
            for n, row in enumerate(b["rows"]):
                cells = table.add_row().cells
                for c, value in enumerate(row + [""] * (width - len(row))):
                    cells[c].text = ""
                    para = cells[c].paragraphs[0]
                    for t, bold, italic, _code in inline_runs(value):
                        run = para.add_run(t)
                        run.bold = True if n == 0 else (bold or None)
                        run.italic = italic or None
            doc.add_paragraph()
        elif kind == "image":
            src = Path(b["path"])
            if base is not None and not src.is_absolute():
                src = base / src
            if src.is_file():
                try:
                    doc.add_picture(str(src), width=Inches(6))
                    if b.get("alt"):
                        cap = doc.add_paragraph()
                        r = cap.add_run(b["alt"])
                        r.italic = True
                except Exception:  # noqa: BLE001 - a picture python-docx cannot read is skipped, said in text
                    doc.add_paragraph("[picture %s could not be added]" % b["path"])
            else:
                doc.add_paragraph("[picture %s not found]" % b["path"])
        elif kind == "hr":
            doc.add_paragraph("―" * 30)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def write_docx(path, markdown: str = "", title: str = "", base: Path | None = None) -> dict:
    docx = _need("docx", "python-docx")
    doc = docx.Document()
    if title:
        doc.core_properties.title = title
        if not str(markdown or "").lstrip().startswith("# "):
            doc.add_heading(title, level=0)
    items = blocks(markdown)
    if title and items and items[0]["type"] == "heading" and items[0]["level"] == 1 and plain(items[0]["text"]) == title:
        doc.add_heading(title, level=0)
        items = items[1:]
    _docx_add_blocks(doc, items, base)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return {"ok": True, "path": str(path), "blocks": len(items)}


def _replace_in_paragraph(p, old: str, new: str) -> int:
    if old not in p.text:
        return 0
    count = 0
    for run in p.runs:   # the easy case keeps the run's formatting
        if old in run.text:
            count += run.text.count(old)
            run.text = run.text.replace(old, new)
    if old in p.text:   # the text spans runs: rebuild it in the first run
        full = p.text
        count += full.count(old)
        runs = p.runs
        if runs:
            runs[0].text = full.replace(old, new)
            for r in runs[1:]:
                r.text = ""
    return count


def edit_docx(path, replace: dict | None = None, append: str = "", out: str | None = None,
              base: Path | None = None) -> dict:
    docx = _need("docx", "python-docx")
    path = Path(path)
    doc = docx.Document(str(path))
    counts = {}
    for old, new in (replace or {}).items():
        n = 0
        for p in doc.paragraphs:
            n += _replace_in_paragraph(p, str(old), str(new))
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        n += _replace_in_paragraph(p, str(old), str(new))
        for section in doc.sections:
            for part in (section.header, section.footer):
                for p in part.paragraphs:
                    n += _replace_in_paragraph(p, str(old), str(new))
        counts[str(old)] = n
    if append:
        _docx_add_blocks(doc, blocks(append), base)
    target = Path(out) if out else path
    doc.save(str(target))
    missing = [k for k, v in counts.items() if not v]
    return {"ok": True, "path": str(target), "replaced": counts, "not_found": missing}


# --------------------------------------------------------------------------- XLSX
def _cell_value(value):
    """'=SUM(..)' stays a formula; numeric strings become numbers; ISO dates become dates."""
    if isinstance(value, str):
        s = value.strip()
        if s.startswith("="):
            return s
        if re.match(r"^-?\d+$", s) and len(s) < 16 and not (s.startswith("0") and len(s) > 1):
            return int(s)
        if re.match(r"^-?\d*\.\d+$", s) or re.match(r"^-?\d+\.\d*$", s):
            try:
                return float(s)
            except ValueError:
                return value
        if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
            try:
                return datetime.date.fromisoformat(s)
            except ValueError:
                return value
    return value


def write_xlsx(path, sheets: list, charts: list | None = None) -> dict:
    """sheets: [{"name", "rows": [[...]], "widths": {"A": 20}, "number_format": {"C": "#,##0.00"},
    "freeze": "A2", "header": true}]; charts: [{"sheet", "type": bar|col|line|pie|scatter|area, "title",
    "data": "B1:C10" (with the header row), "categories": "A2:A10", "anchor": "E2", "x_title", "y_title"}]."""
    openpyxl = _need("openpyxl", "openpyxl")
    from openpyxl.styles import Font, PatternFill  # pylint: disable=import-outside-toplevel
    from openpyxl.utils import get_column_letter  # pylint: disable=import-outside-toplevel
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    formulas = 0
    if not sheets:
        raise WriteError("give at least one sheet: sheets=[{\"name\": \"Data\", \"rows\": [[...], ...]}]")
    for spec in sheets:
        ws = wb.create_sheet(title=str(spec.get("name") or "Sheet%d" % (len(wb.sheetnames) + 1))[:31])
        rows = spec.get("rows") or []
        for row in rows:
            values = [_cell_value(v) for v in (row if isinstance(row, list) else [row])]
            formulas += sum(1 for v in values if isinstance(v, str) and v.startswith("="))
            ws.append(values)
        if rows and spec.get("header", True):
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EEF7")
            ws.freeze_panes = spec.get("freeze") or "A2"
        elif spec.get("freeze"):
            ws.freeze_panes = spec["freeze"]
        for c in range(1, ws.max_column + 1):   # auto width from the content, capped
            letter = get_column_letter(c)
            width = max((len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, min(ws.max_row, 300) + 1)), default=8)
            ws.column_dimensions[letter].width = min(60, max(8, width + 2))
        for letter, width in (spec.get("widths") or {}).items():
            ws.column_dimensions[str(letter)].width = float(width)
        for letter, fmt in (spec.get("number_format") or {}).items():
            for cell in ws[str(letter)][1:]:
                cell.number_format = str(fmt)
        for cell in ws._cells.values():   # noqa: SLF001 - dates get a date format
            if isinstance(cell.value, datetime.date):
                cell.number_format = "yyyy-mm-dd"
    made = [add_chart(wb, c) for c in (charts or [])]
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return {"ok": True, "path": str(path), "sheets": wb.sheetnames, "formulas": formulas, "charts": made}


def add_chart(wb, spec: dict) -> str:
    from openpyxl import chart as xl  # pylint: disable=import-outside-toplevel
    from openpyxl.utils.cell import range_boundaries  # pylint: disable=import-outside-toplevel
    ws = wb[spec["sheet"]] if spec.get("sheet") in wb.sheetnames else wb.worksheets[0]
    kind = str(spec.get("type") or "bar").lower()
    cls = {"bar": xl.BarChart, "col": xl.BarChart, "column": xl.BarChart, "line": xl.LineChart, "pie": xl.PieChart,
           "scatter": xl.ScatterChart, "area": xl.AreaChart}.get(kind)
    if cls is None:
        raise WriteError("chart type must be bar, column, line, pie, scatter or area")
    ch = cls()
    if kind == "bar":
        ch.type = "bar"
    elif kind in ("col", "column"):
        ch.type = "col"
    ch.title = spec.get("title") or None
    data_ref = str(spec.get("data") or "")
    if not data_ref:
        min_col, min_row, max_col, max_row = 2, 1, ws.max_column, ws.max_row
    else:
        min_col, min_row, max_col, max_row = range_boundaries(data_ref)
    if spec.get("categories"):
        c1, r1, c2, r2 = range_boundaries(str(spec["categories"]))
    else:
        c1, r1, c2, r2 = 1, min_row + 1, 1, max_row
    cats = xl.Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
    if kind == "scatter":
        for col in range(min_col, max_col + 1):
            values = xl.Reference(ws, min_col=col, min_row=min_row, max_row=max_row)
            series = xl.Series(values, cats, title_from_data=True)
            series.marker.symbol = "circle"
            ch.series.append(series)
    else:
        data = xl.Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
    if kind != "pie":
        if spec.get("x_title"):
            ch.x_axis.title = spec["x_title"]
        if spec.get("y_title"):
            ch.y_axis.title = spec["y_title"]
    ch.width, ch.height = float(spec.get("width", 16)), float(spec.get("height", 8))
    anchor = spec.get("anchor") or "%s2" % _col_letter(ws.max_column + 2)
    ws.add_chart(ch, anchor)
    return "%s chart on %s at %s" % (kind, ws.title, anchor)


def _col_letter(n: int) -> str:
    out = ""
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out or "A"


def edit_xlsx(path, cells: dict | None = None, append_rows: dict | None = None, charts: list | None = None,
              out: str | None = None) -> dict:
    """cells: {"Sheet1!B2": 5, "C3": "=B2*2"} (no sheet = the first); append_rows: {"Sheet1": [[...], ...]}."""
    openpyxl = _need("openpyxl", "openpyxl")
    path = Path(path)
    wb = openpyxl.load_workbook(str(path))
    changed = 0
    for ref, value in (cells or {}).items():
        sheet, _, addr = str(ref).rpartition("!")
        sheet = sheet.strip("'")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        ws[addr] = _cell_value(value)
        changed += 1
    for sheet, rows in (append_rows or {}).items():
        ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(str(sheet)[:31])
        for row in rows or []:
            ws.append([_cell_value(v) for v in row])
            changed += 1
    made = [add_chart(wb, c) for c in (charts or [])]
    target = Path(out) if out else path
    wb.save(str(target))
    return {"ok": True, "path": str(target), "changed": changed, "charts": made}


# --------------------------------------------------------------------------- PPTX
def write_pptx(path, slides: list, title: str = "", base: Path | None = None, widescreen: bool = True) -> dict:
    """slides: [{"title", "bullets": [...] | "text": "...", "image": "charts/x.png", "notes": "...",
    "table": [[...]], "layout": "title" | "section" | "content"}] - the first slide becomes a title slide
    when `title` is given."""
    pptx = _need("pptx", "python-pptx")
    from pptx.util import Inches, Pt  # pylint: disable=import-outside-toplevel
    prs = pptx.Presentation()
    if widescreen:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    if title:
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = title
        if len(s.placeholders) > 1:
            s.placeholders[1].text = datetime.date.today().strftime("%d %B %Y")
    if not slides and not title:
        raise WriteError("give slides=[{\"title\": ..., \"bullets\": [...]}]")
    for spec in slides or []:
        layout = str(spec.get("layout") or "content").lower()
        if layout == "title":
            s = prs.slides.add_slide(prs.slide_layouts[0])
            s.shapes.title.text = plain(spec.get("title", ""))
            if len(s.placeholders) > 1:
                s.placeholders[1].text = plain(spec.get("text", "") or spec.get("subtitle", ""))
        elif layout == "section":
            s = prs.slides.add_slide(prs.slide_layouts[2])
            s.shapes.title.text = plain(spec.get("title", ""))
            if len(s.placeholders) > 1:
                s.placeholders[1].text = plain(spec.get("text", ""))
        else:
            has_side = bool(spec.get("image") or spec.get("table"))
            s = prs.slides.add_slide(prs.slide_layouts[5 if (has_side and not spec.get("bullets") and not spec.get("text")) else 1])
            s.shapes.title.text = plain(spec.get("title", ""))
            body = None
            for ph in s.placeholders:
                if ph.placeholder_format.idx == 1:
                    body = ph
            items = spec.get("bullets") or ([ln for ln in str(spec.get("text", "")).split("\n") if ln.strip()] if spec.get("text") else [])
            if body is not None:
                if has_side:
                    body.width = int(W * 0.46)
                body.left, body.top = Inches(0.6), Inches(1.5)
                body.height = H - Inches(2.0)
                tf = body.text_frame
                tf.clear()
                for n, item in enumerate(items):
                    level = 0
                    text = str(item)
                    m = re.match(r"^(\s*)[-*]\s+(.*)$", text)
                    if m:
                        level, text = min(4, len(m.group(1)) // 2), m.group(2)
                    para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                    para.level = level
                    for t, bold, italic, _code in inline_runs(text):
                        r = para.add_run()
                        r.text = t
                        r.font.bold = bold or None
                        r.font.italic = italic or None
                        r.font.size = Pt(22 if level == 0 else 18)
            left = int(W * 0.52) if (items and has_side) else Inches(0.8)
            width = (W - left - Inches(0.6))
            if spec.get("image"):
                src = Path(spec["image"])
                if base is not None and not src.is_absolute():
                    src = base / src
                if src.is_file():
                    pic = s.shapes.add_picture(str(src), left, Inches(1.6), width=width)
                    if pic.height > H - Inches(2.0):   # keep it on the slide
                        ratio = (H - Inches(2.0)) / pic.height
                        pic.height, pic.width = int(pic.height * ratio), int(pic.width * ratio)
                    if spec.get("alt"):
                        pic._element.nvPicPr.cNvPr.set("descr", str(spec["alt"]))  # noqa: SLF001
            elif spec.get("table"):
                rows = spec["table"]
                cols = max(len(r) for r in rows)
                shape = s.shapes.add_table(len(rows), cols, left, Inches(1.6), width, Inches(0.4) * len(rows))
                for r, row in enumerate(rows):
                    for c in range(cols):
                        cell = shape.table.cell(r, c)
                        cell.text = plain(str(row[c]) if c < len(row) else "")
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(14)
        if spec.get("notes"):
            s.notes_slide.notes_text_frame.text = str(spec["notes"])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return {"ok": True, "path": str(path), "slides": len(prs.slides)}


def edit_pptx(path, replace: dict | None = None, slides: list | None = None, out: str | None = None,
              base: Path | None = None) -> dict:
    """Replace text in every slide (and notes), and/or append slides written like write_pptx's."""
    pptx = _need("pptx", "python-pptx")
    path = Path(path)
    prs = pptx.Presentation(str(path))
    counts = {str(k): 0 for k in (replace or {})}
    frames = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
        if slide.has_notes_slide:
            frames.append(slide.notes_slide.notes_text_frame)
    for tf in frames:
        for p in tf.paragraphs:
            for old, new in (replace or {}).items():
                counts[str(old)] += _replace_in_paragraph(p, str(old), str(new))
    target = Path(out) if out else path
    prs.save(str(target))
    added = 0
    if slides:
        tmp = target.with_name(target.stem + ".new.pptx")
        write_pptx(tmp, slides, base=base, widescreen=prs.slide_width > prs.slide_height * 1.5)
        extra = pptx.Presentation(str(tmp))
        prs = pptx.Presentation(str(target))
        for s in extra.slides:   # copy text slides over (layout 1 + shapes' XML)
            new = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1])
            for shape in s.shapes:
                if shape.shape_type != 13:   # pictures need their image part; skipped in the copy
                    new.shapes._spTree.insert_element_before(__import__("copy").deepcopy(shape._element), "p:extLst")  # noqa: SLF001
            added += 1
        prs.save(str(target))
        tmp.unlink(missing_ok=True)
    return {"ok": True, "path": str(target), "replaced": counts, "not_found": [k for k, v in counts.items() if not v],
            "added_slides": added}


# --------------------------------------------------------------------------- PDF tools (pypdf)
def parse_ranges(spec: str, count: int) -> list:
    """"1-3,5,8-" -> [[0,1,2],[4],[7..count-1]] (1-based, inclusive, open end)."""
    groups = []
    for part in re.split(r"[,;\s]+", str(spec or "").strip()):
        if not part:
            continue
        m = re.match(r"^(\d*)-(\d*)$", part)
        if m:
            a = int(m.group(1) or 1)
            b = int(m.group(2) or count)
        elif part.isdigit():
            a = b = int(part)
        else:
            raise WriteError("bad page range %r (use e.g. 1-3,5,8-)" % part)
        a, b = max(1, a), min(count, b)
        if a <= b:
            groups.append(list(range(a - 1, b)))
    return groups


def pdf_info(path) -> dict:
    pypdf = _need("pypdf", "pypdf")
    r = pypdf.PdfReader(str(path))
    meta = r.metadata or {}
    fields = r.get_fields() or {}
    return {"ok": True, "pages": len(r.pages), "encrypted": bool(r.is_encrypted), "title": str(meta.get("/Title", "") or ""),
            "author": str(meta.get("/Author", "") or ""), "fields": len(fields),
            "size": [float(r.pages[0].mediabox.width), float(r.pages[0].mediabox.height)] if r.pages else []}


def pdf_merge(inputs: list, out) -> dict:
    pypdf = _need("pypdf", "pypdf")
    if len(inputs or []) < 2:
        raise WriteError("give at least two PDFs to merge")
    w = pypdf.PdfWriter()
    for p in inputs:
        w.append(str(p))
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    with open(out, "wb") as fh:
        w.write(fh)
    return {"ok": True, "path": str(out), "pages": len(w.pages), "inputs": len(inputs)}


def pdf_split(path, out_dir, ranges: str = "") -> dict:
    """Each range becomes its own PDF; no ranges = one PDF per page."""
    pypdf = _need("pypdf", "pypdf")
    r = pypdf.PdfReader(str(path))
    groups = parse_ranges(ranges, len(r.pages)) if ranges else [[i] for i in range(len(r.pages))]
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    made = []
    stem = Path(path).stem
    for g in groups:
        w = pypdf.PdfWriter()
        for i in g:
            w.add_page(r.pages[i])
        name = "%s_p%d%s.pdf" % (stem, g[0] + 1, "" if len(g) == 1 else "-%d" % (g[-1] + 1))
        with open(out_dir / name, "wb") as fh:
            w.write(fh)
        made.append(str(out_dir / name))
    return {"ok": True, "files": made}


def pdf_fields(path) -> dict:
    pypdf = _need("pypdf", "pypdf")
    r = pypdf.PdfReader(str(path))
    out = []
    for name, f in (r.get_fields() or {}).items():
        kind = str(f.get("/FT", ""))
        item = {"name": name, "type": {"/Tx": "text", "/Btn": "checkbox/radio", "/Ch": "choice", "/Sig": "signature"}.get(kind, kind),
                "value": str(f.get("/V", "") or "")}
        states = f.get("/_States_", None)
        if states:
            item["options"] = [str(s) for s in states]
        opts = f.get("/Opt")
        if opts:
            item["options"] = [str(o[1] if isinstance(o, list) else o) for o in opts]
        out.append(item)
    return {"ok": True, "fields": out}


def pdf_fill(path, out, values: dict, flatten: bool = False) -> dict:
    pypdf = _need("pypdf", "pypdf")
    r = pypdf.PdfReader(str(path))
    fields = r.get_fields() or {}
    if not fields:
        raise WriteError("this PDF has no form fields to fill (pdf action=fields lists them)")
    unknown = [k for k in values if k not in fields]
    w = pypdf.PdfWriter(clone_from=r)
    clean = {}
    for k, v in values.items():
        if k not in fields:
            continue
        if str(fields[k].get("/FT")) == "/Btn":
            states = [str(s) for s in (fields[k].get("/_States_") or [])]
            on = next((s for s in states if s != "/Off"), "/Yes")
            clean[k] = on if (v is True or str(v).lower() in ("1", "true", "yes", "on", "x", on.lower().lstrip("/"))) else "/Off"
        else:
            clean[k] = str(v)
    from pypdf.generic import DictionaryObject, NameObject, TextStringObject  # pylint: disable=import-outside-toplevel
    acro = w._root_object.get("/AcroForm")  # noqa: SLF001
    if acro is not None:   # forms made without default resources: give pypdf a font to draw the values
        acro = acro.get_object()
        if "/DA" not in acro:
            acro[NameObject("/DA")] = TextStringObject("/Helv 0 Tf 0 g")
        if "/DR" not in acro:
            helv = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding")})
            acro[NameObject("/DR")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/Helv"): w._add_object(helv)})})  # noqa: SLF001
    for page in w.pages:
        if page.get("/Annots"):
            try:
                w.update_page_form_field_values(page, clean, auto_regenerate=False)
            except Exception:  # noqa: BLE001 - an odd form: set the values directly, the viewer draws them
                for ref in page["/Annots"]:
                    annot = ref.get_object()
                    name = str(annot.get("/T", "") or annot.get("/Parent", {}).get("/T", ""))
                    if name in clean:
                        value = clean[name]
                        target = annot.get("/Parent").get_object() if "/Parent" in annot else annot
                        target[NameObject("/V")] = NameObject(value) if value.startswith("/") else TextStringObject(value)
                        if value.startswith("/"):
                            annot[NameObject("/AS")] = NameObject(value)
    w.set_need_appearances_writer(True)
    if flatten:
        try:
            for page in w.pages:
                w.update_page_form_field_values(page, clean, auto_regenerate=False, flatten=True)
        except TypeError:   # older pypdf: no flatten argument
            pass
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    with open(out, "wb") as fh:
        w.write(fh)
    return {"ok": True, "path": str(out), "filled": sorted(clean), "unknown": unknown}


def pdf_sign_ready(path, out, signers: list | None = None, title: str = "Signatures", date_line: bool = True) -> dict:
    """Append a signature page (one block per signer: a line, the name, a date line) with an empty
    signature FIELD (/FT /Sig) over each line, so Acrobat / Preview / a signing service can sign it there."""
    pypdf = _need("pypdf", "pypdf")
    from pypdf.generic import (ArrayObject, DictionaryObject, FloatObject, NameObject,  # pylint: disable=import-outside-toplevel
                               NumberObject, TextStringObject)
    signers = [str(s).strip() for s in (signers or ["Signature"]) if str(s).strip()][:8] or ["Signature"]
    r = pypdf.PdfReader(str(path))
    w = pypdf.PdfWriter(clone_from=r)
    width = float(r.pages[0].mediabox.width) if r.pages else 595.0
    height = float(r.pages[0].mediabox.height) if r.pages else 842.0
    pdf = SimplePDF(width, height)
    pdf.heading(title, 1)
    pdf.para("Signed by the people below on the date written next to their name.")
    boxes = []
    for name in signers:
        pdf.space(36)
        y = pdf.y
        pdf.line(pdf.margin, y, pdf.margin + 260, y)
        if date_line:
            pdf.line(pdf.margin + 300, y, pdf.margin + 440, y)
        pdf.y -= 4
        pdf.text_at(pdf.margin, pdf.y - 10, name, 10)
        if date_line:
            pdf.text_at(pdf.margin + 300, pdf.y - 10, "Date", 10)
        boxes.append((name, [pdf.margin, y, pdf.margin + 260, y + 34]))
        pdf.y -= 22
    sig_page = pypdf.PdfReader(io.BytesIO(pdf.build())).pages[0]
    w.add_page(sig_page)
    page = w.pages[-1]
    acro = w._root_object.get("/AcroForm")  # noqa: SLF001
    if acro is None:
        acro = DictionaryObject({NameObject("/Fields"): ArrayObject()})
        w._root_object[NameObject("/AcroForm")] = w._add_object(acro)  # noqa: SLF001
    else:
        acro = acro.get_object()
        if "/Fields" not in acro:
            acro[NameObject("/Fields")] = ArrayObject()
    annots = page.get("/Annots")
    if annots is None:
        annots = ArrayObject()
        page[NameObject("/Annots")] = annots
    else:
        annots = annots.get_object()
    used = {str(f.get_object().get("/T", "")) for f in acro["/Fields"]}
    for n, (name, rect) in enumerate(boxes, 1):
        field_name = re.sub(r"[^A-Za-z0-9_]+", "_", name).strip("_") or "Signature%d" % n
        while field_name in used:
            field_name += "_%d" % n
        used.add(field_name)
        field = DictionaryObject({
            NameObject("/FT"): NameObject("/Sig"), NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject("/Widget"), NameObject("/T"): TextStringObject(field_name),
            NameObject("/TU"): TextStringObject("Signature of %s" % name), NameObject("/F"): NumberObject(4),
            NameObject("/Rect"): ArrayObject([FloatObject(v) for v in rect]),
            NameObject("/P"): page.indirect_reference,
        })
        ref = w._add_object(field)  # noqa: SLF001
        annots.append(ref)
        acro["/Fields"].append(ref)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    with open(out, "wb") as fh:
        w.write(fh)
    return {"ok": True, "path": str(out), "pages": len(w.pages), "signature_fields": [b[0] for b in boxes]}


# --------------------------------------------------------------------------- LOKI's own PDF writer
# Standard 14 fonts only (WinAnsi), so it needs nothing installed. Helvetica advance widths (Adobe AFM,
# /1000 em) for 32..126; bold is drawn ~6% wider.
_HELV = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278, 556, 556, 556, 556, 556,
         556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556, 1015, 667, 667, 722, 722, 667, 611, 778, 722, 278,
         500, 667, 556, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469,
         556, 333, 556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500,
         278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584]


def text_width(text: str, size: float, font: str = "F1") -> float:
    if font == "F3":
        return len(text) * 0.6 * size
    total = 0
    for ch in text:
        o = ord(ch)
        total += _HELV[o - 32] if 32 <= o <= 126 else 556
    return total / 1000.0 * size * (1.06 if font == "F2" else 1.0)


def _pdf_str(text: str) -> bytes:
    raw = str(text).encode("cp1252", "replace")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


class SimplePDF:
    """A tiny flowing-layout PDF writer: headings, wrapped paragraphs, lists, code, tables, rules, JPEG
    pictures, page numbers. Fonts: F1 Helvetica, F2 Helvetica-Bold, F3 Courier, F4 Helvetica-Oblique."""

    def __init__(self, width: float = 595.28, height: float = 841.89, margin: float = 56.0) -> None:
        self.w, self.h, self.margin = width, height, margin
        self.pages: list = []
        self.images: list = []   # (name, width_px, height_px, jpeg bytes)
        self.cur: list = []
        self.pages.append(self.cur)
        self.y = self.h - self.margin
        self.title = ""

    def new_page(self) -> None:
        if self.cur:
            self.cur = []
            self.pages.append(self.cur)
        self.y = self.h - self.margin

    def _need_room(self, height: float) -> None:
        if self.y - height < self.margin + 20:
            self.cur = []
            self.pages.append(self.cur)
            self.y = self.h - self.margin

    def space(self, pts: float) -> None:
        self._need_room(pts)
        self.y -= pts

    def text_at(self, x: float, y: float, text: str, size: float = 11, font: str = "F1", gray: float = 0.0) -> None:
        self.cur.append(b"BT %.3f g /%s %.2f Tf %.2f %.2f Td %s Tj ET" % (gray, font.encode(), size, x, y, _pdf_str(text)))

    def line(self, x1, y1, x2, y2, width: float = 0.7, gray: float = 0.3) -> None:
        self.cur.append(b"%.3f G %.2f w %.2f %.2f m %.2f %.2f l S" % (gray, width, x1, y1, x2, y2))

    def rect(self, x, y, w, h, gray: float = 0.95) -> None:
        self.cur.append(b"%.3f g %.2f %.2f %.2f %.2f re f" % (gray, x, y, w, h))

    def wrap(self, runs: list, size: float, width: float) -> list:
        """runs [(text, font)] -> lines [[(text, font)]] that fit `width`."""
        lines, cur, used = [], [], 0.0
        for text, font in runs:
            for word in re.split(r"(\s+)", text):
                if not word:
                    continue
                if word.isspace():
                    if cur:
                        cur.append((" ", font))
                        used += text_width(" ", size, font)
                    continue
                ww = text_width(word, size, font)
                while ww > width and len(word) > 1:   # a word longer than the line is cut
                    cut = max(1, int(len(word) * width / ww) - 1)
                    if cur:
                        lines.append(cur)
                        cur, used = [], 0.0
                    lines.append([(word[:cut], font)])
                    word = word[cut:]
                    ww = text_width(word, size, font)
                if used + ww > width and cur:
                    while cur and cur[-1][0] == " ":
                        cur.pop()
                    lines.append(cur)
                    cur, used = [], 0.0
                cur.append((word, font))
                used += ww
        if cur:
            lines.append(cur)
        return lines

    def runs_of(self, text: str, base: str = "F1") -> list:
        out = []
        for t, bold, italic, code in inline_runs(text):
            out.append((t, "F3" if code else "F2" if bold or base == "F2" else "F4" if italic else base))
        return out

    def flow(self, runs: list, size: float = 11, indent: float = 0.0, lead: float = 1.38, bullet: str = "",
             gray: float = 0.0) -> None:
        width = self.w - 2 * self.margin - indent
        for n, ln in enumerate(self.wrap(runs, size, width)):
            self._need_room(size * lead)
            self.y -= size * lead
            x = self.margin + indent
            if bullet and n == 0:
                self.text_at(x - 12, self.y, bullet, size, "F1", gray)
            parts = []
            for text, font in ln:
                parts.append(b"/%s %.2f Tf %s Tj" % (font.encode(), size, _pdf_str(text)))
            self.cur.append(b"BT %.3f g %.2f %.2f Td " % (gray, x, self.y) + b" ".join(parts) + b" ET")

    def heading(self, text: str, level: int = 1) -> None:
        size = {1: 22, 2: 17, 3: 14, 4: 12.5}.get(level, 12)
        self.space(size * (0.9 if level > 1 else 0.4))
        self._need_room(size * 3)
        self.flow(self.runs_of(plain(text), "F2"), size, lead=1.25, gray=0.08)
        if level <= 2:
            self.y -= 4
            self.line(self.margin, self.y, self.w - self.margin, self.y, 0.6 if level == 1 else 0.3, 0.75)
        self.y -= 4

    def para(self, text: str, size: float = 11) -> None:
        self.flow(self.runs_of(text), size)
        self.y -= 6

    def code(self, text: str) -> None:
        lines = []
        for raw in str(text).split("\n"):
            raw = raw.replace("\t", "    ")
            per = max(10, int((self.w - 2 * self.margin - 16) / (0.6 * 9)))
            lines += [raw[i:i + per] for i in range(0, max(1, len(raw)), per)] or [""]
        for ln in lines:
            self._need_room(12)
            self.rect(self.margin, self.y - 12.5, self.w - 2 * self.margin, 12.5, 0.95)
            self.y -= 12
            self.text_at(self.margin + 8, self.y + 2.5, ln, 9, "F3", 0.15)
        self.y -= 8

    def table(self, rows: list, size: float = 9.5) -> None:
        if not rows:
            return
        cols = max(len(r) for r in rows)
        rows = [r + [""] * (cols - len(r)) for r in rows]
        full = self.w - 2 * self.margin
        natural = [max(20.0, min(full * 0.6, max(text_width(plain(r[c]), size, "F2") for r in rows) + 10)) for c in range(cols)]
        scale = min(1.0, full / sum(natural))
        widths = [n * scale for n in natural]
        if sum(widths) < full * 0.5:
            widths = [w * (full * 0.5 / sum(widths)) for w in widths]
        for n, row in enumerate(rows):
            wrapped = [self.wrap(self.runs_of(cell, "F2" if n == 0 else "F1"), size, max(10, widths[c] - 8)) for c, cell in enumerate(row)]
            height = max(1, max(len(w) for w in wrapped)) * size * 1.3 + 6
            self._need_room(height)
            top = self.y
            if n == 0:
                self.rect(self.margin, top - height, sum(widths), height, 0.92)
            x = self.margin
            for c, lines in enumerate(wrapped):
                ty = top - size - 2
                for ln in lines:
                    parts = [b"/%s %.2f Tf %s Tj" % (f.encode(), size, _pdf_str(t)) for t, f in ln]
                    self.cur.append(b"BT 0 g %.2f %.2f Td " % (x + 4, ty) + b" ".join(parts) + b" ET")
                    ty -= size * 1.3
                x += widths[c]
            self.line(self.margin, top - height, self.margin + sum(widths), top - height, 0.4, 0.7)
            self.y = top - height
        self.y -= 10

    def image_jpeg(self, data: bytes, max_width: float | None = None) -> bool:
        import doc_read  # pylint: disable=import-outside-toplevel
        wpx, hpx = doc_read.image_size(data)
        if not wpx or data[:2] != b"\xff\xd8":
            return False
        width = min(max_width or (self.w - 2 * self.margin), self.w - 2 * self.margin)
        height = width * hpx / wpx
        if height > self.h - 2 * self.margin - 30:
            height = self.h - 2 * self.margin - 30
            width = height * wpx / hpx
        self._need_room(height + 6)
        name = "Im%d" % (len(self.images) + 1)
        self.images.append((name, wpx, hpx, data))
        self.y -= height
        self.cur.append(b"q %.2f 0 0 %.2f %.2f %.2f cm /%s Do Q" % (width, height, self.margin, self.y, name.encode()))
        self.y -= 8
        return True

    def build(self, page_numbers: bool = False) -> bytes:
        pages = [p for p in self.pages if p] or [[]]
        objs: list = []

        def add(body: bytes) -> int:
            objs.append(body)
            return len(objs)
        catalog = add(b"")
        pages_obj = add(b"")
        fonts = {"F1": "Helvetica", "F2": "Helvetica-Bold", "F3": "Courier", "F4": "Helvetica-Oblique"}
        font_refs = {k: add(b"<< /Type /Font /Subtype /Type1 /BaseFont /%s /Encoding /WinAnsiEncoding >>" % v.encode())
                     for k, v in fonts.items()}
        img_refs = {}
        for name, wpx, hpx, data in self.images:
            img_refs[name] = add(b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceRGB "
                                 b"/BitsPerComponent 8 /Filter /DCTDecode /Length %d >>\nstream\n" % (wpx, hpx, len(data))
                                 + data + b"\nendstream")
        res = b"<< /Font << " + b" ".join(b"/%s %d 0 R" % (k.encode(), v) for k, v in font_refs.items()) + b" >>"
        if img_refs:
            res += b" /XObject << " + b" ".join(b"/%s %d 0 R" % (k.encode(), v) for k, v in img_refs.items()) + b" >>"
        res += b" >>"
        kids = []
        for n, ops in enumerate(pages, 1):
            body = b"\n".join(ops)
            if page_numbers and len(pages) > 1:
                label = "%d / %d" % (n, len(pages))
                body += b"\nBT 0.5 g /F1 8.5 Tf %.2f %.2f Td %s Tj ET" % (
                    self.w / 2 - text_width(label, 8.5) / 2, self.margin / 2, _pdf_str(label))
            data = zlib.compress(body)
            content = add(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(data) + data + b"\nendstream")
            kids.append(add(b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 %.2f %.2f] /Resources %s /Contents %d 0 R >>"
                            % (pages_obj, self.w, self.h, res, content)))
        objs[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_obj
        objs[pages_obj - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % k for k in kids), len(kids))
        info = add(b"<< /Producer (LOKI) /Title %s /CreationDate (D:%s) >>" % (
            _pdf_str(self.title or ""), datetime.datetime.now().strftime("%Y%m%d%H%M%S").encode()))
        out = io.BytesIO()
        out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = []
        for n, body in enumerate(objs, 1):
            offsets.append(out.tell())
            out.write(b"%d 0 obj\n" % n + body + b"\nendobj\n")
        xref = out.tell()
        out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
        for off in offsets:
            out.write(b"%010d 00000 n \n" % off)
        out.write(b"trailer\n<< /Size %d /Root %d 0 R /Info %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, catalog, info, xref))
        return out.getvalue()


def _image_bytes_as_jpeg(path: Path) -> bytes:
    data = path.read_bytes()
    if data[:2] == b"\xff\xd8":
        return data
    try:
        from PIL import Image  # pylint: disable=import-outside-toplevel
        img = Image.open(io.BytesIO(data))
        if img.mode in ("RGBA", "LA", "P"):
            bg = Image.new("RGB", img.size, (255, 255, 255))
            img = img.convert("RGBA")
            bg.paste(img, mask=img.split()[-1])
            img = bg
        out = io.BytesIO()
        img.convert("RGB").save(out, "JPEG", quality=90)
        return out.getvalue()
    except Exception:  # noqa: BLE001
        return b""


def markdown_to_pdf_builtin(markdown: str, out, title: str = "", base: Path | None = None, paper: str = "a4") -> dict:
    size = (612.0, 792.0) if str(paper).lower() == "letter" else (595.28, 841.89)
    pdf = SimplePDF(*size)
    pdf.title = title
    items = blocks(markdown)
    if title and not (items and items[0]["type"] == "heading" and items[0]["level"] == 1):
        pdf.heading(title, 1)
    numbers: dict = {}
    for b in items:
        kind = b["type"]
        if kind != "number":
            numbers.clear()
        if kind == "heading":
            pdf.heading(b["text"], b["level"])
        elif kind == "para":
            pdf.para(b["text"])
        elif kind == "bullet":
            pdf.flow(pdf.runs_of(b["text"]), 11, indent=18 + 16 * b["level"], bullet="•" if b["level"] == 0 else "–")
            pdf.y -= 2
        elif kind == "number":
            numbers[b["level"]] = numbers.get(b["level"], 0) + 1
            for deeper in [k for k in numbers if k > b["level"]]:
                numbers.pop(deeper)
            pdf.flow(pdf.runs_of(b["text"]), 11, indent=20 + 16 * b["level"], bullet="%d." % numbers[b["level"]])
            pdf.y -= 2
        elif kind == "quote":
            top = pdf.y
            pdf.flow(pdf.runs_of(b["text"]), 11, indent=16, gray=0.3)
            pdf.line(pdf.margin + 4, top - 2, pdf.margin + 4, pdf.y - 3, 2.0, 0.75)
            pdf.y -= 6
        elif kind == "code":
            pdf.code(b["text"])
        elif kind == "table":
            pdf.table(b["rows"])
        elif kind == "hr":
            pdf.space(6)
            pdf.line(pdf.margin, pdf.y, pdf.w - pdf.margin, pdf.y, 0.5, 0.7)
            pdf.y -= 8
        elif kind == "pagebreak":
            pdf.new_page()
        elif kind == "image":
            src = Path(b["path"])
            if base is not None and not src.is_absolute():
                src = base / src
            data = _image_bytes_as_jpeg(src) if src.is_file() else b""
            if not data or not pdf.image_jpeg(data):
                pdf.para("[picture: %s]" % (b.get("alt") or b["path"]))
            elif b.get("alt"):
                pdf.flow([(b["alt"], "F4")], 9, gray=0.35)
                pdf.y -= 6
    data = pdf.build(page_numbers=True)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    Path(out).write_bytes(data)
    return {"ok": True, "path": str(out), "engine": "built-in", "pages": len([p for p in pdf.pages if p]) or 1}


PDF_CSS = """
@page { size: %(paper)s; margin: 22mm 20mm; }
body { font-family: -apple-system, "Segoe UI", Helvetica, Arial, sans-serif; font-size: 11pt; line-height: 1.5; color: #1d232b; }
h1 { font-size: 22pt; border-bottom: 2px solid #dfe4ea; padding-bottom: 6px; margin: 0 0 14px; }
h2 { font-size: 16pt; border-bottom: 1px solid #e6e9ee; padding-bottom: 4px; margin-top: 26px; }
h3 { font-size: 13pt; margin-top: 20px; }
table { border-collapse: collapse; width: 100%%; margin: 12px 0; font-size: 9.5pt; }
th, td { border: 1px solid #d5dbe3; padding: 5px 8px; text-align: left; vertical-align: top; }
th { background: #eef2f7; }
tr:nth-child(even) td { background: #fafbfc; }
code { font-family: Menlo, Consolas, monospace; font-size: 9.5pt; background: #f2f4f7; padding: 1px 4px; border-radius: 3px; }
pre { background: #f5f7fa; padding: 10px 12px; border-radius: 6px; overflow: hidden; white-space: pre-wrap; }
pre code { background: none; padding: 0; }
blockquote { border-left: 3px solid #c6ced9; margin: 10px 0; padding: 2px 14px; color: #4a5562; }
img { max-width: 100%%; }
hr { border: 0; border-top: 1px solid #dfe4ea; margin: 20px 0; }
.pb { page-break-after: always; }
"""


def markdown_to_html(markdown: str, title: str = "", base: Path | None = None, paper: str = "a4") -> str:
    import base64  # pylint: disable=import-outside-toplevel
    import html as _h  # pylint: disable=import-outside-toplevel

    def inline(text: str) -> str:
        out = []
        for t, bold, italic, code in inline_runs(text):
            t = _h.escape(t)
            out.append("<code>%s</code>" % t if code else "<strong>%s</strong>" % t if bold else "<em>%s</em>" % t if italic else t)
        return "".join(out)
    parts, open_list = [], None
    for b in blocks(markdown):
        kind = b["type"]
        want = {"bullet": "ul", "number": "ol"}.get(kind)
        if open_list and open_list != want:
            parts.append("</%s>" % open_list)
            open_list = None
        if want and not open_list:
            parts.append("<%s>" % want)
            open_list = want
        if kind == "heading":
            parts.append("<h%d>%s</h%d>" % (b["level"], inline(b["text"]), b["level"]))
        elif kind == "para":
            parts.append("<p>%s</p>" % inline(b["text"]))
        elif want:
            parts.append("<li style=\"margin-left:%dpx\">%s</li>" % (b["level"] * 18, inline(b["text"])))
        elif kind == "quote":
            parts.append("<blockquote>%s</blockquote>" % inline(b["text"]))
        elif kind == "code":
            parts.append("<pre><code>%s</code></pre>" % _h.escape(b["text"]))
        elif kind == "table" and b["rows"]:
            head, *body = b["rows"]
            parts.append("<table><thead><tr>%s</tr></thead><tbody>%s</tbody></table>" % (
                "".join("<th>%s</th>" % inline(c) for c in head),
                "".join("<tr>%s</tr>" % "".join("<td>%s</td>" % inline(c) for c in r) for r in body)))
        elif kind == "hr":
            parts.append("<hr>")
        elif kind == "pagebreak":
            parts.append('<div class="pb"></div>')
        elif kind == "image":
            src = Path(b["path"])
            if base is not None and not src.is_absolute():
                src = base / src
            if src.is_file():
                mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif",
                        ".svg": "image/svg+xml", ".webp": "image/webp"}.get(src.suffix.lower(), "image/png")
                parts.append('<figure><img alt="%s" src="data:%s;base64,%s"><figcaption>%s</figcaption></figure>' % (
                    _h.escape(b.get("alt", "")), mime, base64.b64encode(src.read_bytes()).decode(), _h.escape(b.get("alt", ""))))
    if open_list:
        parts.append("</%s>" % open_list)
    head = "<h1>%s</h1>" % _h.escape(title) if title and not str(markdown).lstrip().startswith("# ") else ""
    return ("<!doctype html><html><head><meta charset=\"utf-8\"><title>%s</title><style>%s</style></head><body>%s%s</body></html>"
            % (_h.escape(title or "Document"), PDF_CSS % {"paper": "Letter" if paper == "letter" else "A4"}, head, "\n".join(parts)))


def markdown_to_pdf(markdown: str, out, title: str = "", base: Path | None = None, engine: str = "",
                    paper: str = "a4") -> dict:
    """Chromium (best: real fonts, any language) > reportlab > the built-in writer (always works)."""
    engine = (engine or os.getenv("LOKI_PDF_ENGINE", "") or "auto").lower()
    tried = []
    if engine in ("auto", "chromium"):
        try:
            from playwright.sync_api import sync_playwright  # type: ignore # pylint: disable=import-outside-toplevel
            html = markdown_to_html(markdown, title, base, paper)
            with sync_playwright() as p:
                exe = os.getenv("LOKI_SUITE_CHROMIUM", "")
                browser = p.chromium.launch(**({"executable_path": exe} if exe and os.path.exists(exe) else {}))
                try:
                    page = browser.new_page()
                    page.route("**/*", lambda r: r.abort() if not r.request.url.startswith(("data:", "about:")) else r.continue_())
                    page.set_content(html, wait_until="load")
                    Path(out).parent.mkdir(parents=True, exist_ok=True)
                    page.pdf(path=str(out), format="Letter" if paper == "letter" else "A4", print_background=True,
                             display_header_footer=True, header_template="<span></span>",
                             footer_template='<div style="font-size:8px;width:100%;text-align:center;color:#888">'
                                             '<span class="pageNumber"></span> / <span class="totalPages"></span></div>',
                             margin={"top": "18mm", "bottom": "18mm", "left": "16mm", "right": "16mm"})
                finally:
                    browser.close()
            return {"ok": True, "path": str(out), "engine": "chromium"}
        except Exception as exc:  # noqa: BLE001 - no Playwright/Chromium here: the next engine
            tried.append("chromium: %s" % type(exc).__name__)
            if engine == "chromium":
                raise WriteError("Chromium could not make the PDF (%s) - install the chromium add-on" % exc) from exc
    if engine in ("auto", "reportlab"):
        try:
            return _reportlab_pdf(markdown, out, title, base, paper)
        except ImportError:
            tried.append("reportlab: not installed")
        except Exception as exc:  # noqa: BLE001
            tried.append("reportlab: %s" % type(exc).__name__)
    got = markdown_to_pdf_builtin(markdown, out, title, base, paper)
    if tried:
        got["note"] = "; ".join(tried)
    return got


def _reportlab_pdf(markdown: str, out, title: str, base: Path | None, paper: str) -> dict:
    import html as _h  # pylint: disable=import-outside-toplevel
    from reportlab.lib import colors  # type: ignore # pylint: disable=import-outside-toplevel
    from reportlab.lib.pagesizes import A4, LETTER  # type: ignore # pylint: disable=import-outside-toplevel
    from reportlab.lib.styles import getSampleStyleSheet  # type: ignore # pylint: disable=import-outside-toplevel
    from reportlab.platypus import (Image, PageBreak, Paragraph, Preformatted, SimpleDocTemplate,  # type: ignore # pylint: disable=import-outside-toplevel
                                    Spacer, Table, TableStyle)
    styles = getSampleStyleSheet()

    def inline(text):
        out_ = []
        for t, bold, italic, code in inline_runs(text):
            t = _h.escape(t)
            out_.append('<font face="Courier">%s</font>' % t if code else "<b>%s</b>" % t if bold else "<i>%s</i>" % t if italic else t)
        return "".join(out_)
    story = []
    if title and not str(markdown).lstrip().startswith("# "):
        story.append(Paragraph(_h.escape(title), styles["Title"]))
    counter = 0
    for b in blocks(markdown):
        kind = b["type"]
        counter = counter + 1 if kind == "number" else 0
        if kind == "heading":
            story.append(Paragraph(inline(b["text"]), styles["Heading%d" % min(4, b["level"])]))
        elif kind == "para":
            story.append(Paragraph(inline(b["text"]), styles["BodyText"]))
        elif kind in ("bullet", "number"):
            story.append(Paragraph(inline(b["text"]), styles["BodyText"], bulletText="•" if kind == "bullet" else "%d." % counter))
        elif kind == "quote":
            story.append(Paragraph("<i>%s</i>" % inline(b["text"]), styles["BodyText"]))
        elif kind == "code":
            story.append(Preformatted(b["text"], styles["Code"]))
        elif kind == "table" and b["rows"]:
            data = [[Paragraph(inline(c), styles["BodyText"]) for c in r] for r in b["rows"]]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d5dbe3")),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
            story.append(Spacer(1, 8))
        elif kind == "pagebreak":
            story.append(PageBreak())
        elif kind == "hr":
            story.append(Spacer(1, 10))
        elif kind == "image":
            src = Path(b["path"])
            if base is not None and not src.is_absolute():
                src = base / src
            if src.is_file():
                img = Image(str(src))
                ratio = min(1.0, 440.0 / img.drawWidth)
                img.drawWidth, img.drawHeight = img.drawWidth * ratio, img.drawHeight * ratio
                story.append(img)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out), pagesize=LETTER if paper == "letter" else A4, title=title or "")

    def number(canvas, d):
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(d.pagesize[0] / 2, 20, str(d.page))
    doc.build(story, onFirstPage=number, onLaterPages=number)
    return {"ok": True, "path": str(out), "engine": "reportlab"}
