"""Session 29, offline: documents, files and data. Every sample file is made here.

  reading (doc_read: PDF incl. a scan through OCR, DOCX, XLSX, PPTX, CSV, HTML, EPUB, EML with an attachment,
  pictures, notebooks; anchors, sections, chunks, the cache), writing (doc_write: Word, Excel with formulas
  and charts, PowerPoint, PDF merge/split/fill/sign-ready, Markdown -> PDF), the data sandbox (profile, charts,
  custom code, no network, no writes outside, no subprocess, time limit) and SQL, document collections in the
  one knowledge index + memory.pack citations, the Files API, the agent tools through loki.call_tool (the Files
  view in headless Chromium is tests/test_files_ui.py).

    python tests/test_documents.py
"""
from __future__ import annotations

import asyncio
import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "engine"))
TMP = Path(tempfile.mkdtemp(prefix="loki-docs-"))
DATA = TMP / "data"
WS = TMP / "ws"
WS.mkdir(parents=True, exist_ok=True)
os.environ["LOKI_DATA"] = str(DATA)
os.environ["LOKI_WORKSPACE"] = str(WS)
os.environ.pop("LOKI_JOB_ID", None)
os.environ.pop("LOKI_DOCLING", None)
os.environ["LOKI_PDF_ENGINE"] = "builtin"   # fast; the Chromium path is checked once explicitly (engine="auto")

PASS: list = []
FAIL: list = []


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print(("  PASS " if ok else "  FAIL ") + name + ("" if ok else "  <- %s" % (str(detail)[:400],)))


def has(mod: str) -> bool:
    return importlib.util.find_spec(mod) is not None


def run(coro):
    return asyncio.run(coro)


import data_sandbox  # noqa: E402
import doc_read  # noqa: E402
import doc_write  # noqa: E402
import files_api  # noqa: E402
import knowledge  # noqa: E402

SAMPLES = WS / "samples"
SAMPLES.mkdir()


# --------------------------------------------------------------------------- hand-made OOXML (no libraries)
def make_zip(path: Path, parts: dict) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in parts.items():
            z.writestr(name, text)
    return path


WNS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
DOCX = make_zip(SAMPLES / "handmade.docx", {
    "word/document.xml": '<w:document %s><w:body>'
    '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Lease terms</w:t></w:r></w:p>'
    '<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Notice</w:t></w:r><w:r><w:t xml:space="preserve"> period is sixty days.</w:t></w:r></w:p>'
    '<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr><w:r><w:t>Rent due monthly</w:t></w:r></w:p>'
    '<w:p><w:r><w:br w:type="page"/></w:r><w:hyperlink r:id="rId9"><w:r><w:t>the landlord site</w:t></w:r></w:hyperlink></w:p>'
    '<w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>Payments</w:t></w:r></w:p>'
    '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Month</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Amount</w:t></w:r></w:p></w:tc></w:tr>'
    '<w:tr><w:tc><w:p><w:r><w:t>May</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>1200</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
    '</w:body></w:document>' % WNS,
    "word/styles.xml": '<w:styles %s><w:style w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>'
    '<w:style w:styleId="Heading2"><w:name w:val="heading 2"/></w:style></w:styles>' % WNS,
    "word/numbering.xml": '<w:numbering %s><w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:numFmt w:val="bullet"/></w:lvl>'
    '</w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>' % WNS,
    "word/_rels/document.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId9" Type="http://x/hyperlink" Target="https://example.com/landlord" TargetMode="External"/></Relationships>',
})
SNS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
XLSX = make_zip(SAMPLES / "handmade.xlsx", {
    "xl/workbook.xml": '<workbook %s><sheets><sheet name="Sales" sheetId="1" r:id="rId1"/><sheet name="Notes" sheetId="2" r:id="rId2"/></sheets></workbook>' % SNS,
    "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="x/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="x/worksheet" Target="worksheets/sheet2.xml"/></Relationships>',
    "xl/sharedStrings.xml": '<sst %s><si><t>Region</t></si><si><t>Units</t></si><si><t>Day</t></si><si><t>North</t></si><si><r><t>So</t></r><r><t>uth</t></r></si></sst>' % SNS,
    "xl/styles.xml": '<styleSheet %s><cellXfs><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>' % SNS,
    "xl/worksheets/sheet1.xml": '<worksheet %s><sheetData>'
    '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c></row>'
    '<row r="2"><c r="A2" t="s"><v>3</v></c><c r="B2"><v>12</v></c><c r="C2" s="1"><v>46027</v></c></row>'
    '<row r="3"><c r="A3" t="s"><v>4</v></c><c r="B3"><v>7.5</v></c></row>'
    '<row r="4"><c r="A4" t="inlineStr"><is><t>Total</t></is></c><c r="B4"><f>SUM(B2:B3)</f><v>19.5</v></c></row>'
    '</sheetData></worksheet>' % SNS,
    "xl/worksheets/sheet2.xml": '<worksheet %s><sheetData><row r="1"><c r="A1" t="str"><v>hello</v></c></row></sheetData></worksheet>' % SNS,
})
ANS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
PPTX = make_zip(SAMPLES / "handmade.pptx", {
    "ppt/presentation.xml": '<p:presentation %s><p:sldIdLst><p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId1"/></p:sldIdLst></p:presentation>' % ANS,
    "ppt/_rels/presentation.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="x/slide" Target="slides/slide2.xml"/><Relationship Id="rId2" Type="x/slide" Target="slides/slide1.xml"/></Relationships>',
    "ppt/slides/slide1.xml": '<p:sld %s><p:cSld><p:spTree>'
    '<p:sp><p:nvSpPr><p:cNvPr id="3" name="body"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="2000"/></a:xfrm></p:spPr>'
    '<p:txBody><a:p><a:r><a:t>Second in reading order</a:t></a:r></a:p></p:txBody></p:sp>'
    '<p:sp><p:nvSpPr><p:cNvPr id="2" name="t"/><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="100"/></a:xfrm></p:spPr>'
    '<p:txBody><a:p><a:r><a:t>Roadmap</a:t></a:r></a:p></p:txBody></p:sp>'
    '<p:pic><p:nvPicPr><p:cNvPr id="4" name="pic" descr="A fox logo"/></p:nvPicPr><p:spPr><a:xfrm><a:off x="0" y="5000"/></a:xfrm></p:spPr></p:pic>'
    '</p:spTree></p:cSld></p:sld>' % ANS,
    "ppt/slides/_rels/slide1.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId5" Type="x/notesSlide" Target="../notesSlides/notesSlide1.xml"/></Relationships>',
    "ppt/notesSlides/notesSlide1.xml": '<p:notes %s><p:cSld><p:spTree><p:sp><p:nvSpPr><p:cNvPr id="2" name="n"/><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr>'
    '<p:txBody><a:p><a:r><a:t>Say hello first</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:notes>' % ANS,
    "ppt/slides/slide2.xml": '<p:sld %s><p:cSld><p:spTree><p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="5" name="tbl"/></p:nvGraphicFramePr>'
    '<a:graphic><a:graphicData><a:tbl><a:tr><a:tc><a:txBody><a:p><a:r><a:t>Q</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>Revenue</a:t></a:r></a:p></a:txBody></a:tc></a:tr>'
    '<a:tr><a:tc><a:txBody><a:p><a:r><a:t>Q1</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>40</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl>'
    '</a:graphicData></a:graphic></p:graphicFrame></p:spTree></p:cSld></p:sld>' % ANS,
})
EPUB = make_zip(SAMPLES / "book.epub", {
    "META-INF/container.xml": '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
    "OEBPS/content.opf": '<package xmlns="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata><dc:title>The Fox Book</dc:title></metadata>'
    '<manifest><item id="c1" href="c1.xhtml"/><item id="c2" href="c2.xhtml"/></manifest><spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>',
    "OEBPS/c1.xhtml": "<html><body><h1>Chapter One</h1><p>The fox ran across the snowy island.</p></body></html>",
    "OEBPS/c2.xhtml": "<html><body><h1>Chapter Two</h1><p>The fox found a lighthouse.</p></body></html>",
})

print("reading - hand-made files, standard library")
d = doc_read.convert(DOCX)
check("docx: read", d["ok"], d.get("error"))
check("docx: headings become anchored sections", [s["anchor"] for s in d["sections"]] == ["s1-lease-terms", "s2-payments"], d["sections"])
check("docx: bold, list, link, table", "**Notice** period is sixty days." in d["markdown"] and "- Rent due monthly" in d["markdown"]
      and "[the landlord site](https://example.com/landlord)" in d["markdown"] and "| May" in d["markdown"], d["markdown"])
check("docx: a page break moves the page number", d["sections"][1]["page"] == 2, d["sections"])
check("markdown carries {#anchor} marks; strip_anchors removes them", "{#s2-payments}" in d["markdown"]
      and "{#" not in doc_read.strip_anchors(d["markdown"]))

x = doc_read.convert(XLSX)
check("xlsx: one section per sheet", [s["anchor"] for s in x["sections"]] == ["sheet-sales", "sheet-notes"], x["sections"])
check("xlsx: shared + rich + inline strings, numbers, dates, cached formula values", all(t in x["markdown"] for t in
      ("North", "South", "Total", "19.5", "2026-01-05", "SUM(B2:B3)", "hello")), x["markdown"])
tables = doc_read.xlsx_tables(XLSX)
check("xlsx_tables gives typed rows", tables[0]["rows"][1][:2] == ["North", 12] and tables[0]["formulas"] == {"B4": "SUM(B2:B3)"}, tables[0])

p = doc_read.convert(PPTX)
check("pptx: slides in presentation order", [s["anchor"] for s in p["sections"]] == ["slide-1", "slide-2"], p["sections"])
s1 = p["sections"][0]
check("pptx: title from the title placeholder, shapes top-first, alt text, notes",
      s1["title"] == "Slide 1: Roadmap" and s1["text"].index("Second in reading order") < s1["text"].index("A fox logo")
      and "Say hello first" in s1["text"], s1)
check("pptx: tables", "| Q1" in p["sections"][1]["text"] and "Revenue" in p["sections"][1]["text"], p["sections"][1])

e = doc_read.convert(EPUB)
check("epub: title and chapters in spine order", e["title"] == "The Fox Book" and [s["anchor"] for s in e["sections"]] == ["ch-1", "ch-2"]
      and "lighthouse" in e["sections"][1]["text"], e)

(SAMPLES / "page.html").write_text("<html><head><title>Help page</title></head><body><nav>menu</nav><main><h1>Install</h1>"
                                   "<p>Run setup.</p><h2>Update</h2><p>Click Update everything.</p></main></body></html>")
h = doc_read.convert(SAMPLES / "page.html")
check("html: main content, sections at headings", h["ok"] and any(s["title"] == "Update" for s in h["sections"])
      and "Click Update everything" in h["markdown"], h)

(SAMPLES / "sales.csv").write_text("date,region,units,price\n" + "\n".join(
    "2026-%02d-%02d,%s,%d,%.2f" % (1 + i // 28, 1 + i % 28, ["North", "South", "East"][i % 3], 10 + i, 2.5 + (i % 7) / 10)
    for i in range(60)))
c = doc_read.convert(SAMPLES / "sales.csv")
check("csv: a table with its size", c["ok"] and "60 rows, 4 columns" in c["sections"][0]["title"] and "| North" in c["markdown"], c["sections"][0]["title"])

nb = {"cells": [{"cell_type": "markdown", "source": ["# Notes"]}, {"cell_type": "code", "source": ["1+1"], "outputs": [{"text": ["2"]}]}]}
(SAMPLES / "n.ipynb").write_text(json.dumps(nb))
check("notebook: cells with outputs", "Output:" in doc_read.convert(SAMPLES / "n.ipynb")["markdown"])

import email.message  # noqa: E402
msg = email.message.EmailMessage()
msg["From"], msg["To"], msg["Subject"] = "mia@example.com", "owner@example.com", "Invoice for May"
msg.set_content("Hi,\nthe invoice is attached.\n")
msg.add_attachment(b"item,cost\nhosting,12\n", maintype="text", subtype="csv", filename="invoice.csv")
(SAMPLES / "mail.eml").write_bytes(bytes(msg))
m = doc_read.convert(SAMPLES / "mail.eml")
check("eml: headers, body, attachment listed and read", m["ok"] and "**Subject:** Invoice for May" in m["markdown"]
      and "invoice.csv" in m["markdown"] and any(s["anchor"] == "att-1" and "hosting" in s["text"] for s in m["sections"]), m["markdown"])

check("unknown binary is refused politely", "cannot read" in doc_read.convert(make_zip(SAMPLES / "x.bin", {"a": "b"}))["error"]
      or doc_read.convert(SAMPLES / "x.bin")["ok"])
check("a broken docx gives an error, not an exception", not doc_read.convert(SAMPLES / "page.html").get("error")
      and "zip" in doc_read.convert((SAMPLES / "broken.docx").write_bytes(b"nope") and SAMPLES / "broken.docx")["error"])
bomb = SAMPLES / "entity.docx"
make_zip(bomb, {"word/document.xml": '<!DOCTYPE x [<!ENTITY a "aaaa">]><w:document %s><w:body/></w:document>' % WNS})
check("XML with entity declarations is refused", "entities" in doc_read.convert(bomb)["error"])

check("md_table pads, escapes pipes, drops empty rows", doc_read.md_table([["a", "b|c"], ["", ""], [1, None]]) ==
      "| a   | b\\|c |\n| --- | ---- |\n| 1   |      |")
check("MarkItDown's partial-numbering merge", doc_read.merge_partial_numbering(".1\n\nThe intent\n.2\nNext") == ".1 The intent\n.2 Next")

print("writing - Markdown -> PDF (built-in) and PDF reading")
MD = ("# Annual report\n\nRevenue grew **12%** in *2026*.\n\n## Regions\n\n| Region | Sales |\n| --- | --- |\n| North | 10 |\n"
      "| South | 7 |\n\n- first point\n  - nested point\n1. step one\n2. step two\n\n> a quote\n\n```\ncode()\n```\n\n\\pagebreak\n\n"
      "## Appendix\n\nThe appendix lives on page two. " + "Long text. " * 120)
pdf1 = SAMPLES / "report.pdf"
got = doc_write.markdown_to_pdf_builtin(MD, pdf1, title="Annual report")
check("built-in PDF writer makes a valid PDF", got["ok"] and pdf1.read_bytes().startswith(b"%PDF-1.4") and got["pages"] >= 2, got)
pages = doc_read._pdf_pages_builtin(pdf1)  # noqa: SLF001
check("built-in PDF reader reads LOKI's own PDF in page order", len(pages) == got["pages"] and "Revenue grew" in pages[0]
      and "appendix lives" in pages[1], [p[:80] for p in pages])
r = doc_read.convert(pdf1)
check("convert(pdf): page anchors p1..pN", r["ok"] and [s["anchor"] for s in r["sections"]][:2] == ["p1", "p2"] and "page two" in r["sections"][1]["text"], r.get("error"))
check("select(pages='2') and select(section='p1')", doc_read.select(r, pages="2")[0]["anchor"] == "p2"
      and doc_read.select(r, section="p1")[0]["anchor"] == "p1")
ch = doc_read.chunks(r)
check("chunks keep the page anchor", ch and all(a.startswith("p") for _n, _t, a in ch) and any(a == "p2" for _n, _t, a in ch))
cache = DATA / "documents" / "cache"
first = doc_read.convert_cached(pdf1, cache)
second = doc_read.convert_cached(pdf1, cache)
check("conversion cache hit on the second read", not first.get("cached") and second.get("cached"))
html_out = doc_write.markdown_to_html(MD, "Annual report")
check("Markdown -> HTML (for the Chromium PDF) has tables and lists", "<table>" in html_out and "<ol>" in html_out and "<blockquote>" in html_out)
auto = doc_write.markdown_to_pdf(MD, SAMPLES / "auto.pdf", "Annual report", engine="auto")
check("markdown_to_pdf picks the best engine available (%s)" % auto.get("engine"), auto["ok"] and (SAMPLES / "auto.pdf").read_bytes()[:4] == b"%PDF")

# a scanned PDF: a picture of text, no text layer
if has("PIL"):
    from PIL import Image, ImageDraw, ImageFont  # noqa: E402
    img = Image.new("RGB", (1300, 420), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 64)
    except OSError:
        font = ImageFont.load_default()
    draw.text((40, 60), "INVOICE NUMBER 4471", fill="black", font=font)
    draw.text((40, 200), "TOTAL DUE 980 EUR", fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=95)
    scan = doc_write.SimplePDF()
    scan.image_jpeg(buf.getvalue())
    (SAMPLES / "scan.pdf").write_bytes(scan.build())
    img.save(SAMPLES / "photo.png")
    s = doc_read.convert(SAMPLES / "scan.pdf")
    if doc_read.ocr_engines() and doc_read.can_render_pdf():
        check("scanned PDF: OCR reads the page (%s)" % s["backend"], "4471" in s["markdown"] and "+ocr" in s["backend"], s["markdown"][:300])
    else:
        check("scanned PDF without OCR says so", any("scan" in w for w in s["warnings"]), s["warnings"])
    off = doc_read.convert(SAMPLES / "scan.pdf", ocr="off")
    check("ocr='off' never runs OCR", "+ocr" not in off["backend"])
    pi = doc_read.convert(SAMPLES / "photo.png")
    if doc_read.ocr_engines():
        check("picture: OCR text (%s)" % pi["backend"], "TOTAL" in pi["markdown"].upper(), pi["markdown"][:200])
    else:
        check("picture without OCR: a clear warning", any("OCR" in w for w in pi["warnings"]), pi["warnings"])
    seen = []
    pv = doc_read.convert(SAMPLES / "photo.png", ocr="off", describe=True, vision=lambda png, q: seen.append(len(png)) or "A white card with an invoice total.")
    check("picture: describe=True asks the vision model with PNG bytes", seen and "invoice total" in pv["markdown"] and "+vision" in pv["backend"])
check("image_size reads PNG/JPEG headers", doc_read.image_size(b"\x89PNG\r\n\x1a\n" + b"\0\0\0\rIHDR" + (640).to_bytes(4, "big") + (480).to_bytes(4, "big")) == (640, 480))

print("writing - Office files")
if has("docx"):
    out = doc_write.write_docx(SAMPLES / "made.docx", MD.replace("\\pagebreak", ""), title="Annual report")
    back = doc_read.convert(SAMPLES / "made.docx")
    check("write_docx -> read back: headings, table, numbered and nested lists", out["ok"] and "| North" in back["markdown"]
          and "1. step one" in back["markdown"] and "  - nested point" in back["markdown"] and any(s["title"] == "Regions" for s in back["sections"]),
          back["markdown"][:600])
    ed = doc_write.edit_docx(SAMPLES / "made.docx", {"North": "Nordic", "missing words": "x"}, append="## Added\n\nA new paragraph.")
    back = doc_read.convert(SAMPLES / "made.docx")
    check("edit_docx replaces (in tables too), appends, reports what it did not find", ed["replaced"]["North"] >= 1
          and ed["not_found"] == ["missing words"] and "Nordic" in back["markdown"] and "A new paragraph." in back["markdown"], ed)
else:
    try:
        doc_write.write_docx(SAMPLES / "made.docx", "x")
        said = ""
    except doc_write.WriteError as exc:
        said = str(exc)
    check("python-docx missing: write_docx says how to install it", "pip install python-docx" in said, said)
if has("openpyxl"):
    import openpyxl  # noqa: E402
    out = doc_write.write_xlsx(SAMPLES / "made.xlsx", [{"name": "Sales", "rows": [["Region", "Q1", "Q2", "Total"], ["North", "10", 12, "=B2+C2"],
                                                                               ["South", 7, "9.5", "=B3+C3"], ["When", "2026-03-01", "", ""]],
                                                       "number_format": {"D": "#,##0.00"}}],
                               charts=[{"sheet": "Sales", "type": "bar", "title": "By quarter", "data": "B1:C3", "categories": "A2:A3"},
                                       {"sheet": "Sales", "type": "line", "data": "B1:B3", "anchor": "H20"}])
    wb = openpyxl.load_workbook(SAMPLES / "made.xlsx")
    ws = wb["Sales"]
    check("write_xlsx: real formulas, numbers typed, date typed, header bold, frozen", ws["D2"].value == "=B2+C2" and ws["B2"].value == 10
          and ws["C3"].value == 9.5 and ws["B4"].value.year == 2026 and ws["A1"].font.bold and ws.freeze_panes == "A2", out)
    check("write_xlsx: two charts", len(ws._charts) == 2 and out["formulas"] == 2, out)  # noqa: SLF001
    ed = doc_write.edit_xlsx(SAMPLES / "made.xlsx", {"Sales!B2": 11, "E1": "=SUM(B2:C3)"}, {"Extra": [["a", 1]]})
    wb = openpyxl.load_workbook(SAMPLES / "made.xlsx")
    check("edit_xlsx: cells, a new sheet's rows", wb["Sales"]["B2"].value == 11 and wb["Sales"]["E1"].value == "=SUM(B2:C3)"
          and wb["Extra"]["B1"].value == 1, ed)
if has("pptx"):
    (SAMPLES / "chart.png").write_bytes((SAMPLES / "photo.png").read_bytes() if (SAMPLES / "photo.png").exists() else b"")
    slides = [{"title": "Plan", "bullets": ["Ship it", "  - test first"], "notes": "Keep it short"},
              {"title": "Numbers", "table": [["Q", "Rev"], ["Q1", "40"]]}]
    if (SAMPLES / "photo.png").exists():
        slides.append({"title": "Picture", "image": "samples/photo.png", "alt": "the scan"})
    out = doc_write.write_pptx(SAMPLES / "made.pptx", slides, title="Deck", base=WS)
    back = doc_read.convert(SAMPLES / "made.pptx")
    check("write_pptx -> read back: title slide, bullets, notes, table%s" % (", picture alt" if len(slides) == 3 else ""),
          out["slides"] == len(slides) + 1 and "Ship it" in back["markdown"] and "Keep it short" in back["markdown"]
          and "| Q1" in back["markdown"] and (len(slides) == 2 or "the scan" in back["markdown"]), back["markdown"][:500])
    ed = doc_write.edit_pptx(SAMPLES / "made.pptx", {"Ship it": "Launch it"}, slides=[{"title": "Extra", "bullets": ["one more"]}])
    back = doc_read.convert(SAMPLES / "made.pptx")
    check("edit_pptx: replace text and add a slide", "Launch it" in back["markdown"] and "one more" in back["markdown"] and ed["added_slides"] == 1, ed)

print("PDF tools")
if has("pypdf"):
    import pypdf  # noqa: E402
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, NumberObject, TextStringObject  # noqa: E402
    merged = doc_write.pdf_merge([pdf1, SAMPLES / "auto.pdf"], SAMPLES / "merged.pdf")
    check("pdf_merge", merged["pages"] == len(pypdf.PdfReader(pdf1).pages) + len(pypdf.PdfReader(SAMPLES / "auto.pdf").pages), merged)
    sp = doc_write.pdf_split(SAMPLES / "merged.pdf", SAMPLES / "split", "1,2-3")
    check("pdf_split by ranges", len(sp["files"]) == 2 and len(pypdf.PdfReader(sp["files"][1]).pages) == 2, sp)
    check("parse_ranges: open end and clamping", doc_write.parse_ranges("2-,9", 4) == [[1, 2, 3]])
    # a form with one text field and one checkbox, made by hand
    w = pypdf.PdfWriter(clone_from=pypdf.PdfReader(pdf1))
    page = w.pages[0]
    fields = ArrayObject()
    for name, ft, rect in (("FullName", "/Tx", [72, 700, 300, 720]), ("Agree", "/Btn", [72, 660, 90, 678])):
        f = DictionaryObject({NameObject("/FT"): NameObject(ft), NameObject("/T"): TextStringObject(name),
                              NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
                              NameObject("/Rect"): ArrayObject([FloatObject(v) for v in rect]), NameObject("/F"): NumberObject(4),
                              NameObject("/P"): page.indirect_reference})
        if ft == "/Btn":
            f[NameObject("/AP")] = DictionaryObject({NameObject("/N"): DictionaryObject({NameObject("/Yes"): w._add_object(DictionaryObject()),  # noqa: SLF001
                                                                                         NameObject("/Off"): w._add_object(DictionaryObject())})})  # noqa: SLF001
        ref = w._add_object(f)  # noqa: SLF001
        page[NameObject("/Annots")] = page.get("/Annots", ArrayObject())
        page["/Annots"].append(ref)
        fields.append(ref)
    w._root_object[NameObject("/AcroForm")] = w._add_object(DictionaryObject({NameObject("/Fields"): fields}))  # noqa: SLF001
    with open(SAMPLES / "form.pdf", "wb") as fh:
        w.write(fh)
    ff = doc_write.pdf_fields(SAMPLES / "form.pdf")
    check("pdf_fields lists text and checkbox fields", {f["name"]: f["type"] for f in ff["fields"]} == {"FullName": "text", "Agree": "checkbox/radio"}, ff)
    fill = doc_write.pdf_fill(SAMPLES / "form.pdf", SAMPLES / "filled.pdf", {"FullName": "Mia Lund", "Agree": True, "Nope": 1})
    vals = pypdf.PdfReader(SAMPLES / "filled.pdf").get_fields()
    check("pdf_fill sets text and ticks the box; unknown names reported", vals["FullName"].get("/V") == "Mia Lund"
          and str(vals["Agree"].get("/V")) == "/Yes" and fill["unknown"] == ["Nope"], (fill, {k: v.get("/V") for k, v in vals.items()}))
    sig = doc_write.pdf_sign_ready(SAMPLES / "form.pdf", SAMPLES / "to_sign.pdf", ["Ann Lee", "Bo Chan"])
    rd = pypdf.PdfReader(SAMPLES / "to_sign.pdf")
    sigs = [k for k, v in (rd.get_fields() or {}).items() if v.get("/FT") == "/Sig"]
    check("pdf_sign_ready: a signature page + one /Sig field per signer, the form kept", len(rd.pages) == len(pypdf.PdfReader(SAMPLES / "form.pdf").pages) + 1
          and sorted(sigs) == ["Ann_Lee", "Bo_Chan"] and "FullName" in (rd.get_fields() or {}) and "Ann Lee" in (rd.pages[-1].extract_text() or ""), (sig, sigs))
    check("pdf_info", doc_write.pdf_info(SAMPLES / "form.pdf")["fields"] == 2)

print("data sandbox + SQL")
if data_sandbox.has_pandas():
    a = data_sandbox.analyze([SAMPLES / "sales.csv"], workspace=WS)
    check("profile: finding with totals, top group, trend", a["ok"] and "units: total" in a["finding"] and "Largest units by region" in a["finding"]
          and "from 2026-01-01" in a["finding"], a)
    check("profile: charts saved as PNG in analysis/<time>/", a["charts"] and all((WS / c).read_bytes()[:4] == b"\x89PNG" for c in a["charts"])
          and all(c.startswith("analysis/") for c in a["charts"]), a["charts"])
    b = data_sandbox.analyze([SAMPLES / "sales.csv"], code="show(df.groupby('region').units.sum())\nfinding = 'East leads'\n"
                             "df.plot(x='date', y='units'); chart('units over time')", workspace=WS)
    check("custom code: show() table, chart(), finding", b["ok"] and b["finding"] == "East leads" and "| East" in b["tables"][0]
          and b["charts"][0].endswith("units_over_time.png"), b)
    attempts = {"socket": "import socket; socket.create_connection(('example.com', 80))",
                "urllib": "import urllib.request; urllib.request.urlopen('http://example.com')",
                "write": "open(%r, 'w').write('x')" % str(TMP / "outside.txt"),
                "subprocess": "import subprocess; subprocess.run(['echo', 'hi'])",
                "system": "import os; os.system('echo hi')",
                "remove": "import os; os.remove(%r)" % str(SAMPLES / "sales.csv")}
    probe = "said = {}\n" + "".join("try:\n    %s\n    said[%r] = 'ran'\nexcept Exception as e:\n    said[%r] = str(e)\n" % (c, k, k)
                                    for k, c in attempts.items()) + "import json; finding = json.dumps(said)"
    res = data_sandbox.analyze([SAMPLES / "sales.csv"], code=probe, workspace=WS, timeout=30)
    said = json.loads(res.get("finding") or "{}")
    for k in attempts:
        check("sandbox refuses %s" % k, "blocked" in said.get(k, "") or "only" in said.get(k, ""), (said.get(k), res.get("error")))
    check("...and nothing was written or deleted outside", not (TMP / "outside.txt").exists() and (SAMPLES / "sales.csv").exists())
    t = data_sandbox.analyze([SAMPLES / "sales.csv"], code="while True: pass", workspace=WS, timeout=2)
    check("time limit stops a runaway analysis", not t["ok"] and "longer than 2s" in t["error"], t)
    mem = data_sandbox.analyze([SAMPLES / "sales.csv"], code="x = bytearray(3 * 1024 ** 3); x[::4096] = b'1' * len(x[::4096])",
                               workspace=WS, timeout=30, memory_mb=400)
    check("memory limit stops a hungry analysis", not mem["ok"] and ("memory" in mem["error"].lower() or "MemoryError" in mem["error"]), mem)
    if has("openpyxl"):
        xa = data_sandbox.analyze([SAMPLES / "made.xlsx"], workspace=WS, charts=False)
        check("analyze reads Excel too", xa["ok"] and "rows x" in xa["finding"], xa)
st = data_sandbox.profile_stdlib([SAMPLES / "sales.csv"], WS)
check("standard-library profile without pandas", st["ok"] and "units: total 2370" in st["finding"] and "Most common region" in st["finding"], st["finding"])
q = data_sandbox.sql("SELECT region, SUM(units) AS u FROM sales GROUP BY region ORDER BY u DESC", [SAMPLES / "sales.csv"], engine="sqlite")
want = {}
for i in range(60):
    want[["North", "South", "East"][i % 3]] = want.get(["North", "South", "East"][i % 3], 0) + 10 + i
check("SQL over CSV (sqlite): right totals, integers stay integers", q["ok"] and q["rows"] == sorted([[k, v] for k, v in want.items()], key=lambda r: -r[1])
      and isinstance(q["rows"][0][1], int) and q["engine"] == "sqlite", q)
q2 = data_sandbox.sql("select h.Region, h.Units from handmade h where h.Units > 10 and h.Region <> 'Total'", [XLSX], engine="sqlite")
check("SQL over XLSX", q2["ok"] and q2["rows"] == [["North", 12]], q2)
for bad in ("DELETE FROM sales", "select 1; drop table sales", "ATTACH 'x.db' AS x", "select * from sales where 1=1 or load_extension('x')",
            "WITH a AS (SELECT 1) INSERT INTO sales VALUES (1)"):
    res = data_sandbox.sql(bad, [SAMPLES / "sales.csv"])
    check("SQL refuses %r" % bad[:30], not res["ok"], res)
slow = data_sandbox.sql("WITH RECURSIVE c(x) AS (SELECT 1 UNION ALL SELECT x+1 FROM c) SELECT count(*) FROM c", [SAMPLES / "sales.csv"],
                        timeout=1.0, engine="sqlite")
check("SQL time limit", not slow["ok"] and "interrupt" in slow["error"].lower(), slow)
if data_sandbox.has_duckdb():
    dq = data_sandbox.sql("SELECT region, SUM(units) AS u FROM sales GROUP BY region ORDER BY u DESC", [SAMPLES / "sales.csv"])
    check("SQL with DuckDB, external access off", dq["ok"] and dq["engine"] == "duckdb" and dq["rows"][0][0] == "East", dq)
check("table names from file names", data_sandbox.table_name(Path("2026 Sales-Q1.csv"), set()) == "t_2026_sales_q1")

print("document collections in the one index + memory.pack citations")
COLL = TMP / "contracts"
COLL.mkdir()
shutil.copy(pdf1, COLL / "report.pdf")
shutil.copy(DOCX, COLL / "lease.docx")
(COLL / "notes.md").write_text("# Notes\n\nThe fox mascot is orange.\n")
(COLL / ".secret").mkdir()
add = knowledge.add_collection(DATA, COLL, "Contracts")
check("add a collection", add["ok"] and knowledge.list_collections(DATA)[0]["name"] == "Contracts", add)
check("adding the home folder is refused", not knowledge.add_collection(DATA, Path.home())["ok"])
idx = knowledge.collection_index(DATA, knowledge.list_collections(DATA)[0])
stat = idx.refresh(budget_s=60)
check("the collection is indexed by the same ProjectIndex (documents counted)", stat["files"] == 3 and idx.status()["documents"] == 2, stat)
hits = knowledge.search_documents(DATA, "notice period sixty days", fresh=False)
check("search cites collection/file#anchor", hits and hits[0]["cite"] == "Contracts/lease.docx#s1-lease-terms", hits[:2])
hits = knowledge.search_documents(DATA, "appendix lives page two", fresh=False)
check("a PDF hit cites the page and gives page=", hits and hits[0]["cite"] == "Contracts/report.pdf#p2" and hits[0]["page"] == 2, hits[:1])
(COLL / "lease.docx").unlink()
idx.refresh(budget_s=30)
check("a removed file drops out of the index", not any("lease" in h["cite"] for h in knowledge.search_documents(DATA, "notice period sixty days", fresh=False)))
import memory  # noqa: E402
pk = memory.pack("what does the appendix say on page two", root=DATA, workspace=WS, lessons=False)
check("memory.pack carries the document passage with its citation", "Contracts/report.pdf#p2" in pk["text"] and "Contracts/report.pdf#p2" in pk["sources"], pk)
pidx = knowledge.project_index(WS, DATA)
pidx.refresh(budget_s=60)
ph = pidx.search("Lease terms notice period sixty days")
check("the workspace's own documents are in the project index too (docx#anchor)", any(h["cite"].startswith("samples/handmade.docx#") for h in ph), ph[:2])
check("remove a collection", knowledge.remove_collection(DATA, "Contracts")["ok"] and not knowledge.list_collections(DATA))
knowledge.add_collection(DATA, COLL, "Contracts")


class FakeAgent:
    WORKSPACE = WS
    DATA_DIR = DATA


class FakeState:
    agent = FakeAgent()


print("Files API")
S = FakeState()
code, lst = files_api.handle(S, "GET", "/files/list?dir=samples", {})
names = [e["name"] for e in lst.get("entries", [])]
check("list: folders first, kinds, sizes", code == 200 and "report.pdf" in names and next(e for e in lst["entries"] if e["name"] == "report.pdf")["kind"] == "pdf", lst)
(WS / ".hidden").mkdir(exist_ok=True)
(WS / "harness_data").mkdir(exist_ok=True)
_c, top = files_api.handle(S, "GET", "/files/list", {})
check("list hides dot folders and harness_data", not any(e["name"] in (".hidden", "harness_data") for e in top["entries"]))
for bad in ("/files/list?dir=../", "/files/preview?path=../../etc/passwd", "/files/preview?path=.hidden/x", "/files/list?dir=harness_data"):
    c2, r2 = files_api.handle(S, "GET", bad, {})
    check("refused: %s" % bad, c2 == 400 and not r2["ok"], r2)
_c, pv = files_api.handle(S, "GET", "/files/preview?path=samples/report.pdf&page=2", {})
check("preview pdf: page 2 text%s" % (" + picture" if doc_read.can_render_pdf() else " (no renderer: text + note)"), pv["kind"] == "pdf" and pv["page"] == 2
      and "page two" in pv["text"] and (pv["image"].startswith("data:image/png;base64,") if doc_read.can_render_pdf() else bool(pv["note"])), {k: str(v)[:80] for k, v in pv.items()})
_c, pv = files_api.handle(S, "GET", "/files/preview?path=samples/handmade.xlsx&sheet=Notes", {})
check("preview xlsx: a sheet as a table", pv["kind"] == "table" and pv["sheets"] == ["Sales", "Notes"] and pv["columns"] == ["hello"], pv)
_c, pv = files_api.handle(S, "GET", "/files/preview?path=samples/sales.csv", {})
check("preview csv", pv["kind"] == "table" and pv["total"] == 60 and len(pv["rows"]) == 60)
_c, pv = files_api.handle(S, "GET", "/files/preview?path=samples/handmade.pptx", {})
check("preview pptx as Markdown without anchor marks", pv["kind"] == "markdown" and "Roadmap" in pv["markdown"] and "{#" not in pv["markdown"], pv)
if (SAMPLES / "photo.png").exists():
    _c, pv = files_api.handle(S, "GET", "/files/preview?path=samples/photo.png", {})
    check("preview image as a data URL", pv["kind"] == "image" and pv["data_url"].startswith("data:image/png;base64,") and pv["width"] == 1300)
_c, rec = files_api.handle(S, "GET", "/files/recent", {})
check("recent files remember previews, newest first", rec["files"] and rec["files"][0]["path"] in ("samples/photo.png", "samples/handmade.pptx"), rec)
c3, up = files_api.save_upload(S, "name=..%2F..%2Fevil%20report.pdf", io.BytesIO(pdf1.read_bytes()), pdf1.stat().st_size)
check("upload: copied into uploads/, name made safe", c3 == 200 and up["path"] == "uploads/evil report.pdf" and (WS / up["path"]).read_bytes() == pdf1.read_bytes(), up)
c4, up2 = files_api.save_upload(S, "name=evil%20report.pdf", io.BytesIO(b"%PDF-1.4 x"), 10)
check("upload: an existing name gets (2)", up2["path"] == "uploads/evil report (2).pdf", up2)
c5, up3 = files_api.save_upload(S, "name=a.txt&dir=../outside", io.BytesIO(b"x"), 1)
check("upload: a folder outside the workspace is refused", c5 == 400, up3)
c6, up4 = files_api.save_upload(S, "name=cut.txt", io.BytesIO(b"abc"), 10)
check("upload: a cut-off body leaves no file", c6 == 400 and not (WS / "uploads" / "cut.txt").exists() and not (WS / "uploads" / "cut.txt.part").exists())
check("safe_name", files_api.safe_name("..\\..\\CON.txt") == "_CON.txt" and files_api.safe_name(".env") == "env")
opened = []
(WS / "tool.exe").write_bytes(b"MZ")
try:
    files_api.open_default(S, "tool.exe", run=opened.append)
    refused = False
except files_api.FilesError:
    refused = True
check("open refuses programs", refused and not opened)
if sys.platform != "win32":
    files_api.open_default(S, "samples/report.pdf", run=opened.append)
    check("open uses the system opener with an argument list", opened and opened[0][-1].endswith("report.pdf") and opened[0][0] in ("open", "xdg-open"), opened)
_c, conv = files_api.handle(S, "POST", "/files/convert", {"path": "samples/handmade.docx"})
check("convert: saves Markdown next to it", conv["ok"] and (WS / "samples/handmade.md").read_text().startswith("# handmade"), conv)
_c, cl = files_api.handle(S, "GET", "/files/collections", {})
check("collections listed with index status", cl["collections"] and cl["collections"][0]["name"] == "Contracts" and cl["collections"][0]["exists"], cl)
_c, cl2 = files_api.handle(S, "GET", "/files/list?root=Contracts", {})
check("browse a collection's folder (hidden folders still hidden)", {e["name"] for e in cl2["entries"]} == {"report.pdf", "notes.md"}, cl2)
_c, srch = files_api.handle(S, "POST", "/files/collections", {"action": "search", "query": "fox mascot orange", "name": "Contracts"})
check("collections search through the API", srch["ok"] and srch["hits"] and srch["hits"][0]["cite"].startswith("Contracts/notes.md"), srch)

print("the agent tools (loki.call_tool)")
import loki  # noqa: E402
for name in ("read_document", "write_document", "edit_document", "pdf_tools", "analyze_data", "data_sql", "documents"):
    check("tool %s is offered with a schema in the documents category" % name, name in loki.MODEL_TOOL_NAMES and name in loki._tool_params()  # noqa: SLF001
          and name in loki.TOOL_CATEGORIES["documents"])
check("write tools count as modifying (verification, Plan mode)", {"write_document", "edit_document", "pdf_tools", "analyze_data"} <= loki.MODIFYING)
res = run(loki.call_tool("read_file", {"filename": "samples/report.pdf"}))
check("read_file on a PDF reads it as a document (no decode error)", "samples/report.pdf - pdf" in res and "Revenue grew" in res, res[:200])
res = run(loki.call_tool("read_document", {"path": "samples/report.pdf", "pages": "2"}))
check("read_document pages='2' + citation hint", "{#p2}" in res and "{#p1}" not in res and "report.pdf#<anchor>" in res, res[:300])
res = run(loki.call_tool("read_document", {"path": "Contracts/notes.md"}))
check("read_document reaches a collection file", "fox mascot" in res, res[:200])
res = run(loki.call_tool("read_document", {"path": "../../etc/passwd"}))
check("read_document refuses paths outside", res.startswith("error"), res)
big = "# Big\n\n" + "\n\n".join("## Part %d\n\n%s" % (i, "word " * 400) for i in range(30))
(WS / "big.md").write_text(big)
res = run(loki.call_tool("read_document", {"path": "big.md"}))
check("a long document comes in parts with how to read on", "Showing" in res and "section=" in res and len(res) < 14000, res[-300:])
if has("docx"):
    res = run(loki.call_tool("write_document", {"path": "out/letter.docx", "title": "Letter", "content": "Dear [NAME],\n\nThanks."}))
    check("write_document .docx", res.startswith("made out/letter.docx") and (WS / "out/letter.docx").is_file(), res)
    res = run(loki.call_tool("write_document", {"path": "out/letter.docx", "content": "x"}))
    check("write_document will not overwrite by accident", "already exists" in res, res)
    res = run(loki.call_tool("edit_document", {"path": "out/letter.docx", "replace": {"[NAME]": "Mia"}}))
    check("edit_document .docx", "edited" in res and "Dear Mia" in doc_read.convert(WS / "out/letter.docx")["markdown"], res)
if has("openpyxl"):
    res = run(loki.call_tool("write_document", {"path": "out/budget.xlsx", "sheets": [{"name": "Budget", "rows": [["Item", "Cost"], ["Rent", 900], ["Food", 300], ["Total", "=SUM(B2:B3)"]]}],
                                                "charts": [{"sheet": "Budget", "type": "pie", "data": "B1:B3", "categories": "A2:A3"}]}))
    check("write_document .xlsx with a formula and a chart", "1 formulas, 1 charts" in res, res)
    res = run(loki.call_tool("write_document", {"path": "out/table.xlsx", "content": "| a | b |\n|---|---|\n| 1 | 2 |"}))
    check("write_document .xlsx from a Markdown table", res.startswith("made") and doc_read.xlsx_tables(WS / "out/table.xlsx")[0]["rows"] == [["a", "b"], [1, 2]], res)
if has("pptx"):
    res = run(loki.call_tool("write_document", {"path": "out/deck.pptx", "title": "Plan", "content": "## Why\n\n- fast\n- local\n\n## How\n\nOne step at a time."}))
    check("write_document .pptx from Markdown (## = slide)", "3 slides" in res, res)
res = run(loki.call_tool("write_document", {"path": "out/notes.pdf", "title": "Notes", "content": "Hello **PDF**."}))
check("write_document .pdf", res.startswith("made out/notes.pdf") and (WS / "out/notes.pdf").read_bytes()[:4] == b"%PDF", res)
res = run(loki.call_tool("write_document", {"path": "out/x.exe", "content": "x"}))
check("write_document refuses other types", res.startswith("error"), res)
if has("pypdf"):
    res = run(loki.call_tool("pdf_tools", {"action": "fields", "path": "samples/form.pdf"}))
    check("pdf_tools fields", "FullName" in res, res)
    res = run(loki.call_tool("pdf_tools", {"action": "fill", "path": "samples/form.pdf", "values": {"FullName": "Bo"}}))
    check("pdf_tools fill writes next to the source", "samples/form_filled.pdf" in res and (WS / "samples/form_filled.pdf").is_file(), res)
    res = run(loki.call_tool("pdf_tools", {"action": "merge", "paths": ["samples/report.pdf", "out/notes.pdf"], "out": "out/all.pdf"}))
    check("pdf_tools merge", res.startswith("made out/all.pdf"), res)
    res = run(loki.call_tool("pdf_tools", {"action": "sign_ready", "path": "out/notes.pdf", "signers": ["Mohamed"]}))
    check("pdf_tools sign_ready", "signature fields for Mohamed" in res and (WS / "out/notes_to_sign.pdf").is_file(), res)
    res = run(loki.call_tool("pdf_tools", {"action": "split", "path": "out/all.pdf"}))
    check("pdf_tools split", res.startswith("made") and "out/all_split/" in res, res)
res = run(loki.call_tool("analyze_data", {"path": "samples/sales.csv"}))
check("analyze_data answers with a finding%s" % (" and chart images" if data_sandbox.has_pandas() else ""), "units: total" in res
      and (("![" in res and "analysis/" in res) if data_sandbox.has_pandas() else True), res[:400])
res = run(loki.call_tool("data_sql", {"query": "SELECT COUNT(*) AS n FROM sales", "files": ["samples/sales.csv"]}))
check("data_sql", "| 60" in res, res)
res = run(loki.call_tool("data_sql", {"query": "SELECT nope FROM sales", "files": ["samples/sales.csv"]}))
check("data_sql errors list the tables and columns", res.startswith("error") and "sales(date, region, units, price)" in res, res)
res = run(loki.call_tool("documents", {"action": "search", "query": "fox mascot orange"}))
check("documents search returns cited passages", "[Contracts/notes.md:1]" in res and "cite" in res, res[:300])
res = run(loki.call_tool("documents", {"action": "search", "query": "appendix lives on page two"}))
check("documents search cites a PDF page", "[Contracts/report.pdf#p2]" in res, res[:300])
res = run(loki.call_tool("documents", {"action": "collections"}))
check("documents collections", "Contracts" in res, res)

print("docs, add-ons, licences")
import addons_install  # noqa: E402
import manifests  # noqa: E402
for n in ("documents", "ocr", "data-analysis", "docling"):
    mf = json.loads((REPO / "capabilities" / n / "manifest.json").read_text())
    check("add-on %s: valid manifest, pip install plan" % n, not manifests.validate(mf) and addons_install.pip_specs(n), manifests.validate(mf))
req = (REPO / "engine" / "requirements.txt").read_text()
check("the small Office/PDF libraries are in requirements.txt", all(x in req for x in ("python-docx", "openpyxl", "python-pptx", "pypdf")))
third = (REPO / "docs" / "THIRD_PARTY.md").read_text()
check("THIRD_PARTY.md lists MarkItDown and the session-29 add-ons", "MarkItDown" in third and "doc_read.py" in third and "RapidOCR" in third)
check("MarkItDown's licence is kept", (REPO / "engine" / "licenses" / "MarkItDown-MIT.txt").is_file())
ship = (REPO / "tools" / "ship.py").read_text()
check("test_documents is in ship.py's SUITES", '"test_documents"' in ship)


print("\n%d passed, %d failed" % (len(PASS), len(FAIL)))
shutil.rmtree(TMP, ignore_errors=True) if not FAIL and not os.getenv("LOKI_SUITE_KEEP") else print("kept " + str(TMP))
if FAIL:
    print("FAILED: " + ", ".join(FAIL))
    sys.exit(1)
